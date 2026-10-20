"""Lee Google Sheets y los archivos de pasarelas de pago y arma los datos del dashboard."""
import os
import io
import json
import re
from datetime import datetime, timedelta

import pandas as pd
import gspread
from google.oauth2.service_account import Credentials

from .config import (ROOT_DIR, CONFIG_FILE, SEED_CONFIG_FILE, CREDENTIALS_FILE, CACHE_FILE, CACHE_VERSION,
                     DEFAULT_CONFIG)
from .utils import (clean_str, fix_mojibake, normalize_col, has_word, parse_date,
                    extract_sheet_id, find_col, best_filled_col)


class BitelitDataLoader:
    def __init__(self, base_dir=ROOT_DIR):
        self.base_dir = base_dir
        self.cache_file = CACHE_FILE
        os.makedirs(os.path.dirname(CONFIG_FILE), exist_ok=True)
        os.makedirs(os.path.dirname(CACHE_FILE), exist_ok=True)
        self.data = {}
        self.gs_client = None
        self.headers_found = {}
        self.source_errors = {}
        self._df_cache = {}
        self._dayfirst = {}
        self.config = self.load_config()

    # ------------------------------------------------------------------ CONFIG
    def load_config(self):
        cfg = json.loads(json.dumps(DEFAULT_CONFIG))
        # Si aún no existe la copia editable (p. ej. volumen nuevo en el servidor), se parte de la del código.
        src = CONFIG_FILE if os.path.exists(CONFIG_FILE) else SEED_CONFIG_FILE
        if os.path.exists(src):
            try:
                with open(src, 'r', encoding='utf-8') as f:
                    saved = json.load(f)
                for k, v in saved.items():
                    if k in cfg and isinstance(v, dict):
                        cfg[k].update(v)
            except Exception as e:
                print(f"Error leyendo sheets_config.json: {e}. Se usan valores por defecto.")
        if not os.path.exists(CONFIG_FILE):
            try:
                self._write_config(cfg)
            except Exception as e:
                print(f"No se pudo crear sheets_config.json: {e}")
        return cfg

    def _write_config(self, cfg):
        with open(CONFIG_FILE, 'w', encoding='utf-8') as f:
            json.dump(cfg, f, ensure_ascii=False, indent=2)

    def get_config(self):
        self.config = self.load_config()
        return self.config

    def update_config(self, new_values):
        """new_values: {fuente: {sheet_id, pestanas, columnas}}"""
        self.config = self.load_config()
        for key, vals in (new_values or {}).items():
            if key not in self.config or not isinstance(vals, dict):
                continue
            if 'sheet_id' in vals:
                sid = extract_sheet_id(vals['sheet_id'])
                if sid or key == 'finanzas':
                    self.config[key]['sheet_id'] = sid
            if 'pestanas' in vals:
                tabs = vals['pestanas']
                if isinstance(tabs, str):
                    tabs = tabs.split(',')
                tabs = [t for t in (clean_str(x) for x in tabs) if t]
                if tabs or key == 'finanzas':
                    self.config[key]['pestanas'] = tabs
            if isinstance(vals.get('columnas'), dict):
                self.config[key].setdefault('columnas', {})
                for ck, cv in vals['columnas'].items():
                    self.config[key]['columnas'][ck] = clean_str(cv)
        self._write_config(self.config)
        return self.config

    def _sid(self, key):
        return self.config[key]['sheet_id']

    def _tabs(self, key):
        return self.config[key].get('pestanas', [])

    def _col_override(self, key, col):
        return clean_str(self.config[key].get('columnas', {}).get(col, ''))

    # ------------------------------------------------------------------ LOAD
    def load_all_data(self, force_reload=False):
        if not force_reload and os.path.exists(self.cache_file):
            try:
                with open(self.cache_file, 'r', encoding='utf-8') as f:
                    cached = json.load(f)
                if cached.get('cache_version') == CACHE_VERSION:
                    self.data = cached
                    print("Loaded data from cache.")
                    return self.data
                print("Caché de una versión anterior: se vuelve a leer todo.")
            except Exception as e:
                print(f"Cache load error: {e}, re-generating...")

        print("Leyendo Google Sheets...")
        self.config = self.load_config()
        self.headers_found = {}
        self.source_errors = {}
        self._df_cache = {}
        tickets_data = self._load_tickets()
        fallas_altan = self._load_fallas_altan()
        procesos_data = self._load_procesos()
        portabilidades_data = self._load_portabilidades()
        port_outs_data = self._load_port_outs()
        activaciones_data = self._load_activaciones()
        finanzas_data = self._load_finanzas()
        self._df_cache = {}

        self.data = {
            'cache_version': CACHE_VERSION,
            'generated_at': datetime.now().isoformat(),
            'overview': self._build_overview(tickets_data, procesos_data, activaciones_data, finanzas_data),
            'tickets': tickets_data,
            'fallas_altan': fallas_altan,
            'procesos': procesos_data,
            'portabilidades': portabilidades_data,
            'port_outs': port_outs_data,
            'activaciones': activaciones_data,
            'finanzas': finanzas_data,
            'fuentes': {
                'config': self.config,
                'columnas': self.headers_found,
                'errores': self.source_errors
            }
        }

        try:
            with open(self.cache_file, 'w', encoding='utf-8') as f:
                json.dump(self.data, f, ensure_ascii=False, indent=2)
            print("Successfully cached dashboard data.")
        except Exception as e:
            print(f"Error saving cache: {e}")

        return self.data

    def _get_gs_client(self):
        if self.gs_client:
            return self.gs_client
        scopes = ['https://www.googleapis.com/auth/spreadsheets.readonly', 'https://www.googleapis.com/auth/drive.readonly']
        # En el servidor: el JSON completo de la cuenta de servicio en la variable GOOGLE_CREDENTIALS_JSON
        creds_env = os.environ.get('GOOGLE_CREDENTIALS_JSON', '').strip()
        if creds_env:
            try:
                credentials = Credentials.from_service_account_info(json.loads(creds_env), scopes=scopes)
                self.gs_client = gspread.authorize(credentials)
                return self.gs_client
            except Exception as e:
                print(f"Error autenticando Google Sheets con GOOGLE_CREDENTIALS_JSON: {e}")
                return None
        creds_path = CREDENTIALS_FILE
        if not os.path.exists(creds_path):  # compatibilidad: ubicación anterior en la raíz
            creds_path = os.path.join(self.base_dir, 'credentials.json')
        if not os.path.exists(creds_path):
            print("No se encontró config/credentials.json. Se requieren credenciales para cargar Google Sheets.")
            return None
        try:
            scopes = ['https://www.googleapis.com/auth/spreadsheets.readonly', 'https://www.googleapis.com/auth/drive.readonly']
            credentials = Credentials.from_service_account_file(creds_path, scopes=scopes)
            self.gs_client = gspread.authorize(credentials)
            return self.gs_client
        except Exception as e:
            print(f"Error autenticando Google Sheets: {e}")
            return None

    def _read_sheet(self, key, tab):
        """Lee una pestaña con get_all_values() (tolera encabezados repetidos o vacíos)
        y regresa un DataFrame de texto. Los duplicados se renombran 'Columna (2)', 'Columna (3)'..."""
        cache_key = (key, self._sid(key), tab)
        if cache_key in self._df_cache:
            return self._df_cache[cache_key]

        client = self._get_gs_client()
        if not client:
            raise RuntimeError('Sin credenciales de Google Sheets (credentials.json)')
        try:
            doc = client.open_by_key(self._sid(key))
            locale = str(getattr(doc, 'locale', '') or '')
            # en_US => mes/día/año ; es_MX y casi todos los demás => día/mes/año
            self._dayfirst[key] = not locale.startswith('en')
            ws = doc.worksheet(tab)
            values = ws.get_all_values()
        except gspread.exceptions.WorksheetNotFound:
            raise
        except Exception as e:
            # Los archivos .xlsx guardados en Drive no se pueden leer con la API de Sheets:
            # se descargan con la API de Drive y se leen como Excel.
            msg = str(e).lower()
            if 'not supported' in msg or 'no es compatible' in msg or 'failed_precondition' in msg or '400' in msg:
                print(f"[{key}] El archivo parece ser .xlsx; se descarga desde Drive...")
                values = self._read_xlsx_from_drive(key, tab)
            else:
                raise
        if not values:
            self.headers_found.setdefault(key, {})[tab] = []
            df = pd.DataFrame()
            self._df_cache[cache_key] = df
            return df

        headers, seen = [], {}
        for i, h in enumerate(values[0]):
            name = clean_str(h) or f'Columna {i + 1}'
            if name in seen:
                seen[name] += 1
                name = f'{name} ({seen[name]})'
            else:
                seen[name] = 1
            headers.append(name)

        width = len(headers)
        rows = [(list(r) + [''] * width)[:width] for r in values[1:]]
        rows = [r for r in rows if any(str(x).strip() for x in r)]
        self.headers_found.setdefault(key, {})[tab] = headers
        print(f"[{key} / {tab}] {len(rows)} filas leídas.")
        df = pd.DataFrame(rows, columns=headers)
        self._df_cache[cache_key] = df
        return df

    def _read_xlsx_from_drive(self, key, tab):
        """Descarga el archivo (Excel) desde Google Drive y regresa sus valores como lista de filas de texto."""
        from openpyxl import load_workbook
        client = self._get_gs_client()
        url = f'https://www.googleapis.com/drive/v3/files/{self._sid(key)}'
        resp = client.http_client.request('get', url, params={'alt': 'media', 'supportsAllDrives': 'true'})
        wb = load_workbook(io.BytesIO(resp.content), read_only=True, data_only=True)
        if tab in wb.sheetnames:
            ws = wb[tab]
        else:
            ws = wb[wb.sheetnames[0]]
            aviso = f"La pestaña '{tab}' no existe en el Excel; se usó '{wb.sheetnames[0]}'. Pestañas disponibles: {', '.join(wb.sheetnames)}"
            self.source_errors.setdefault(key + '_avisos', []).append(aviso)
            print(f"[{key}] {aviso}")

        def fmt(v):
            if v is None:
                return ''
            if isinstance(v, datetime):
                return v.strftime('%Y-%m-%d %H:%M:%S') if (v.hour or v.minute or v.second) else v.strftime('%Y-%m-%d')
            if isinstance(v, float) and v.is_integer():
                return str(int(v))
            return str(v)

        values = [[fmt(c) for c in row] for row in ws.iter_rows(values_only=True)]
        wb.close()
        # quitar filas vacías al inicio (a veces el Excel trae espacio antes del encabezado)
        while values and not any(x.strip() for x in values[0]):
            values.pop(0)
        return values

    def _pd(self, key, val):
        return parse_date(val, self._dayfirst.get(key, True))

    def _register_error(self, key, e):
        msg = str(e) or e.__class__.__name__
        if isinstance(e, gspread.exceptions.WorksheetNotFound):
            msg = f'No existe la pestaña: {msg}'
        elif isinstance(e, gspread.exceptions.SpreadsheetNotFound):
            msg = 'No se encontró el archivo o la cuenta de servicio no tiene acceso (compártelo con el correo de credentials.json).'
        elif isinstance(e, gspread.exceptions.APIError):
            msg = f'Error de la API de Google: {msg}'
        errs = self.source_errors.setdefault(key, [])
        if msg not in errs:
            errs.append(msg)
        print(f"Error en fuente '{key}': {msg}")

    # ------------------------------------------------------------------ TICKETS
    def _ticket_columns(self, df):
        cols = list(df.columns)
        used = set()

        def pick(name, **kw):
            c = find_col(cols, used=used, override=self._col_override('tickets', name), **kw)
            if c:
                used.add(c)
            return c

        m = {}
        m['remedy'] = pick('remedy', include=['remedy', 'numero de ticket', 'no de ticket', 'folio', 'ticket'], exclude=['fecha', 'captura'])
        m['cerrado_fecha'] = pick('cerrado_fecha', include=['cierre', 'cerrad', 'fecha de solucion', 'solucionado'])
        m['fecha'] = pick('fecha', include=['marca temporal', 'fecha'], exclude=['desde cuando'])
        m['respuesta'] = pick('respuesta', include=['respuesta', 'comentario de cierre', 'solucion', 'diagnostico'])
        m['descripcion'] = pick('descripcion', include=['descrip', 'describe', 'detalle', 'problema', 'comentario'], exclude=['captura'])
        m['agente'] = pick('agente', include=['agente', 'ejecutivo', 'asesor', 'atendio', 'direccion de correo', 'correo electronico'])
        m['dn'] = pick('dn', include=['msisdn afectado', 'msisdn', 'numero afectado', 'linea afectada'], words=['dn'], exclude=['captura'])
        m['contacto'] = pick('contacto', include=['contacto'])
        m['desc_alt1'] = pick('desc_alt1', include=['problema exacto'])
        m['desc_alt2'] = pick('desc_alt2', include=['texto para altan'])
        m['tipo'] = pick('tipo', include=['tipo de falla', 'tipo falla', 'tipo de reporte', 'tipo de incidencia', 'categoria'])
        m['estatus'] = pick('estatus', include=['estatus', 'status', 'estado del ticket'])
        return m

    def _load_tickets(self):
        tickets_list = []
        for tab in self._tabs('tickets'):
            try:
                df = self._read_sheet('tickets', tab)
                if df.empty:
                    continue
                m = self._ticket_columns(df)
                print(f"[tickets] columnas usadas: {m}")
                seen_remedy = set()

                def g(row, k):
                    c = m.get(k)
                    return clean_str(row.get(c, '')) if c else ''

                for _, row in df.iterrows():
                    remedy = g(row, 'remedy')
                    if remedy and remedy in seen_remedy:
                        continue
                    if remedy:
                        seen_remedy.add(remedy)

                    desc = g(row, 'descripcion') or g(row, 'desc_alt1') or g(row, 'desc_alt2')
                    tipo = g(row, 'tipo')
                    fecha_str = self._pd('tickets', g(row, 'fecha'))
                    if not desc and not remedy and not tipo:
                        continue

                    agente = g(row, 'agente')
                    if not agente or agente.lower() in ['nan', 'none']:
                        agente = 'Mesa de Control'
                    else:
                        agente = agente.replace('\n', ' ').strip()
                        if '@' in agente:
                            agente = agente.split('@')[0].replace('.', ' ').replace('_', ' ')
                        agente = agente.split(' ')[0].capitalize()

                    texto = f"{tipo} {desc}".lower()
                    motivo = tipo.rstrip('. ').strip()
                    if not motivo:
                        motivo = "Falla de Servicio / Red"
                        if "roaming" in texto or "t-mobile" in texto or re.search(r'\busa\b', texto):
                            motivo = "Roaming Internacional"
                        elif "volte" in texto or "llamada" in texto:
                            motivo = "Llamadas / VoLTE"
                        elif "apn" in texto or "datos" in texto or "navegar" in texto:
                            motivo = "Datos Móviles / APN"
                        elif "esim" in texto or re.search(r'\bsim\b', texto):
                            motivo = "eSIM / SIM Física"
                        elif "portabilidad" in texto:
                            motivo = "Portabilidad"
                        elif "homolog" in texto or "iphone" in texto or "imei" in texto:
                            motivo = "Compatibilidad IMEI / Dispositivo"
                        elif "sms" in texto or "banco" in texto:
                            motivo = "Recepción SMS Bancarios"

                    resp_val = g(row, 'respuesta')
                    resp_l = resp_val.lower()
                    estatus_val = g(row, 'estatus').lower()
                    cerrado = bool(g(row, 'cerrado_fecha')) \
                        or any(x in estatus_val for x in ['cerrad', 'resuelt', 'solucion', 'concluid']) \
                        or any(x in resp_l for x in ['después', 'con base', 'se regularizan', 'resuelto', 'solucionado'])

                    tickets_list.append({
                        'id': remedy or f"TCK-{len(tickets_list) + 1000}",
                        'periodo': 'Actual',
                        'fecha': fecha_str,
                        'dn': g(row, 'dn') or 'N/A',
                        'contacto': g(row, 'contacto'),
                        'agente': agente,
                        'motivo': motivo,
                        'descripcion': desc[:220] + ('...' if len(desc) > 220 else ''),
                        'descripcion_full': desc,
                        'respuesta': resp_val or "Ticket en validación técnica con agregador / Altan Redes",
                        'estatus': "Cerrado" if cerrado else "En Proceso"
                    })
            except Exception as e:
                self._register_error('tickets', e)

        tickets_list.sort(key=lambda x: x.get('fecha', ''), reverse=True)
        total = len(tickets_list)
        cerrados = sum(1 for t in tickets_list if t['estatus'] == 'Cerrado')

        agente_counts, motivo_counts = {}, {}
        for t in tickets_list:
            ag = t['agente']
            if ag and ag not in ['Mesa', 'Mesa de Control', 'None', 'Nan']:
                agente_counts[ag] = agente_counts.get(ag, 0) + 1
            motivo_counts[t['motivo']] = motivo_counts.get(t['motivo'], 0) + 1

        return {
            'total': total,
            'cerrados': cerrados,
            'en_proceso': total - cerrados,
            'tasa_resolucion': round((cerrados / total * 100) if total > 0 else 0, 1),
            'top_agentes': sorted([{'name': k, 'count': v} for k, v in agente_counts.items()], key=lambda x: x['count'], reverse=True)[:6],
            'motivos': sorted([{'name': k, 'count': v} for k, v in motivo_counts.items()], key=lambda x: x['count'], reverse=True),
            'records': tickets_list
        }

    def _load_fallas_altan(self):
        fallas = []
        for tab in self._tabs('tickets'):
            try:
                df = self._read_sheet('tickets', tab)  # usa la copia ya leída
                if df.empty:
                    continue
                m = self._ticket_columns(df)
                for _, r in df.iterrows():
                    tck = clean_str(r.get(m['remedy'], '')) if m.get('remedy') else ''
                    if not tck:
                        continue
                    fallas.append({
                        'ticket': tck,
                        'fecha': self._pd('tickets', r.get(m['fecha'], '')) if m.get('fecha') else '',
                        'agente': clean_str(r.get(m['agente'], '')) if m.get('agente') else '',
                        'dn_afectado': (clean_str(r.get(m['dn'], '')) if m.get('dn') else '') or 'N/A',
                        'tipo_falla': (clean_str(r.get(m['tipo'], '')) if m.get('tipo') else '') or 'Voz / Datos',
                        'problema': clean_str(r.get(m['descripcion'], '')) if m.get('descripcion') else ''
                    })
            except Exception as e:
                self._register_error('tickets', e)
        return {'total': len(fallas), 'records': sorted(fallas, key=lambda x: str(x.get('fecha', '')), reverse=True)}

    # ------------------------------------------------------------------ PROCESOS
    def _load_procesos(self):
        procesos_list, sim_swap_detalles, imei_bloqueos = [], [], []
        tabs = self._tabs('procesos')
        main_tab = tabs[0] if tabs else 'PROCESOS'
        sim_tab = tabs[1] if len(tabs) > 1 else None

        try:
            df_proc = self._read_sheet('procesos', main_tab)
            col_map = {}
            for c in df_proc.columns:
                cn = normalize_col(c)
                if 'proceso' in cn and 'proceso' not in col_map:
                    col_map['proceso'] = c
                elif 'fecha' in cn and 'fecha' not in col_map:
                    col_map['fecha'] = c
                elif 'motivo' in cn and 'motivo' not in col_map:
                    col_map['motivo'] = c
                elif ('recuperar' in cn or ('nuevo' in cn and 'dn' in cn)) and 'dn' not in col_map:
                    col_map['dn'] = c
                elif 'marca' in cn and 'marca' not in col_map:
                    col_map['marca'] = c
                elif 'agente' in cn and 'agente' not in col_map:
                    col_map['agente'] = c
                elif 'comentario' in cn and 'comentario' not in col_map:
                    col_map['comentario'] = c

            def g(row, k):
                c = col_map.get(k)
                return clean_str(row.get(c, '')) if c else ''

            for _, row in df_proc.iterrows():
                proc_tipo = g(row, 'proceso')
                if not proc_tipo:
                    continue
                procesos_list.append({
                    'fecha': self._pd('procesos', g(row, 'fecha')),
                    'proceso': proc_tipo,
                    'motivo': g(row, 'motivo') or 'Gestión de usuario',
                    'dn': g(row, 'dn') or 'N/A',
                    'marca': g(row, 'marca') or 'Bitelit MVNO',
                    'agente': g(row, 'agente') or 'Mesa de Control',
                    'comentario': g(row, 'comentario') or 'Trámite concluido'
                })
        except Exception as e:
            self._register_error('procesos', e)

        if sim_tab:
            try:
                df_sim = self._read_sheet('procesos', sim_tab)
                for _, r in df_sim.iterrows():
                    vals = list(r.values)
                    sim_swap_detalles.append({
                        'fecha': self._pd('procesos', vals[0]) if len(vals) > 0 else '',
                        'dn': clean_str(vals[2]) if len(vals) > 2 else 'N/A',
                        'iccid_nuevo': clean_str(vals[3]) if len(vals) > 3 else 'N/A',
                        'motivo': clean_str(vals[4]) if len(vals) > 4 else 'Extravío / Daño',
                        'solicitante': clean_str(vals[6]) if len(vals) > 6 else 'Cliente'
                    })
            except Exception as e:
                self._register_error('procesos', e)

        proc_counts, motivo_counts, agente_counts = {}, {}, {}
        for p in procesos_list:
            proc_counts[p['proceso']] = proc_counts.get(p['proceso'], 0) + 1
            if p['motivo'] and p['motivo'] != 'Gestión de usuario':
                motivo_counts[p['motivo']] = motivo_counts.get(p['motivo'], 0) + 1
            if p['agente'] and p['agente'] not in ['Mesa de Control', 'nan']:
                agente_counts[p['agente']] = agente_counts.get(p['agente'], 0) + 1

        return {
            'total': len(procesos_list),
            'sim_swap_count': proc_counts.get('Sim Swap', 0),
            'cambio_dn_count': proc_counts.get('Cambio de DN', 0),
            'rollback_count': proc_counts.get('Rollback', 0),
            'actualizacion_datos_count': proc_counts.get('Actualización de datos', 0),
            'by_type': sorted([{'name': k, 'count': v} for k, v in proc_counts.items()], key=lambda x: x['count'], reverse=True),
            'by_motivo': sorted([{'name': k, 'count': v} for k, v in motivo_counts.items()], key=lambda x: x['count'], reverse=True)[:8],
            'by_agente': sorted([{'name': k, 'count': v} for k, v in agente_counts.items()], key=lambda x: x['count'], reverse=True)[:6],
            'sim_swaps_recientes': sim_swap_detalles[:25],
            'bloqueos_imei': imei_bloqueos,
            'records': sorted(procesos_list, key=lambda x: str(x.get('fecha', '')), reverse=True)
        }

    # ------------------------------------------------------------------ PORT IN
    def _portin_operator_col(self, df):
        """Columna del operador de ORIGEN (donador). Prioriza nombres explícitos; si hay varias,
        usa la que tenga más datos. Nunca toma columnas de receptor/destino."""
        override = self._col_override('portabilidades', 'operador')
        if override:
            c = find_col(list(df.columns), override=override)
            if c:
                return c
            print(f"[portabilidades] La columna configurada '{override}' no existe; se intenta detección automática.")

        cols = list(df.columns)
        bloqueo = ['receptor', 'receptora', 'destino', 'fecha', 'nip']
        prioridades = [
            ['donador', 'donante', 'origen', 'procedencia', 'anterior'],
            ['operador', 'compania', 'proveedor', 'carrier'],
        ]
        for keys in prioridades:
            cands = [c for c in cols if any(k in normalize_col(c) for k in keys)
                     and not any(b in normalize_col(c) for b in bloqueo)]
            best = best_filled_col(df, cands)
            if best:
                return best
        # Abreviatura común "IDA" (identificador del donador)
        cands = [c for c in cols if has_word(normalize_col(c), 'ida')]
        return best_filled_col(df, cands)

    def _load_portabilidades(self):
        portabilidades_list = []
        total_hoja = 0
        cols_usadas = {}
        for tab in self._tabs('portabilidades'):
            try:
                df = self._read_sheet('portabilidades', tab)
                if df.empty:
                    continue
                cols = list(df.columns)
                total_hoja += len(df)
                ov = lambda k: self._col_override('portabilidades', k)

                fvc_sol_col = find_col(cols, include=['fvc solicitada', 'fvc solicita', 'solicitada'], override=ov('fvc_solicitada'))
                date_col = find_col(cols, override=ov('fecha')) if ov('fecha') else None
                date_col = date_col or next((c for c in cols if normalize_col(c) == 'fvc'), None) \
                    or find_col(cols, include=['fecha'])
                dn_col = find_col(cols, include=['msisdn a portar', 'msisdn', 'numero a portar', 'numero'], words=['dn'],
                                  exclude=['nip', 'folio', 'contacto'], override=ov('dn'))
                brand_col = find_col(cols, include=['marca'], override=ov('marca'))
                op_col = self._portin_operator_col(df)
                agent_col = find_col(cols, include=['agente', 'ejecutivo', 'asesor', 'capturo', 'vendedor'], override=ov('agente'))

                cols_usadas = {'fecha': date_col, 'fvc_solicitada': fvc_sol_col, 'dn': dn_col,
                               'marca': brand_col, 'operador': op_col, 'agente': agent_col}
                print(f"[portabilidades] columnas usadas: {cols_usadas}")

                for _, r in df.iterrows():
                    # Solo portabilidades con FVC Solicitada = SI
                    if fvc_sol_col:
                        if normalize_col(r.get(fvc_sol_col, '')) not in ('si', 's'):
                            continue
                    dn = clean_str(r.get(dn_col, '')) if dn_col else ''
                    if not dn:
                        continue
                    portabilidades_list.append({
                        'fecha': self._pd('portabilidades', r.get(date_col, '')) if date_col else '',
                        'dn': dn,
                        'marca': (clean_str(r.get(brand_col, '')) if brand_col else '') or 'Sin marca',
                        'operador': (clean_str(r.get(op_col, '')) if op_col else '') or 'Sin dato'
                    })
            except Exception as e:
                self._register_error('portabilidades', e)

        return {
            'total': len(portabilidades_list),
            'total_hoja': total_hoja,
            'columnas_usadas': cols_usadas,
            'records': sorted(portabilidades_list, key=lambda x: str(x.get('fecha', '')), reverse=True)
        }

    # ------------------------------------------------------------------ PORT OUT
    def _load_port_outs(self):
        portouts_list = []
        for tab in self._tabs('port_outs'):
            try:
                df = self._read_sheet('port_outs', tab)
                if df.empty:
                    continue
                cols = list(df.columns)
                ov = lambda k: self._col_override('port_outs', k)
                # 'Fecha' = día en que la línea se va (NO la fecha de activación)
                date_col = find_col(cols, override=ov('fecha')) if ov('fecha') else None
                date_col = date_col or next((c for c in cols if normalize_col(c) == 'fecha'), None) \
                    or find_col(cols, include=['fecha'], exclude=['activ', 'alta'])
                dias_col = find_col(cols, override=ov('dias')) if ov('dias') else None
                dias_col = dias_col or next((c for c in cols if normalize_col(c) == 'dias'), None) \
                    or find_col(cols, include=['dias'])
                dn_col = find_col(cols, include=['msisdn', 'numero'], words=['dn'], override=ov('dn'))
                brand_col = find_col(cols, include=['marca'], override=ov('marca'))
                op_col = find_col(cols, include=['operador', 'receptor', 'destino'], override=ov('operador'))
                self._portout_cols = {'fecha': date_col, 'dias': dias_col, 'dn': dn_col, 'marca': brand_col, 'operador': op_col}
                print(f"[port_outs] columnas usadas: {self._portout_cols}")

                for _, r in df.iterrows():
                    dn = clean_str(r.get(dn_col, '')) if dn_col else ''
                    if not dn:
                        continue
                    portouts_list.append({
                        'fecha': self._pd('port_outs', r.get(date_col, '')) if date_col else '',
                        'dn': dn,
                        'marca': (clean_str(r.get(brand_col, '')) if brand_col else '') or 'Sin marca',
                        'operador': (clean_str(r.get(op_col, '')) if op_col else '') or 'Sin dato',
                        'dias': self._to_number(r.get(dias_col, '')) if dias_col else None
                    })
            except Exception as e:
                self._register_error('port_outs', e)
        con_dias = [p['dias'] for p in portouts_list if p.get('dias') is not None]
        return {
            'total': len(portouts_list),
            'promedio_dias': round(sum(con_dias) / len(con_dias), 1) if con_dias else None,
            'columnas_usadas': getattr(self, '_portout_cols', {}),
            'records': sorted(portouts_list, key=lambda x: str(x.get('fecha', '')), reverse=True)
        }

    @staticmethod
    def _to_number(v):
        s = clean_str(v).replace(',', '')
        if not s:
            return None
        try:
            return float(s)
        except ValueError:
            m = re.search(r'-?\d+(\.\d+)?', s)
            return float(m.group(0)) if m else None

    # ------------------------------------------------------------------ ACTIVACIONES
    def _load_activaciones(self):
        activaciones_list = []
        for sheet_name in self._tabs('activaciones'):
            try:
                df = self._read_sheet('activaciones', sheet_name)
                if df.empty:
                    continue
                col_map = {}
                for c in df.columns:
                    cn = normalize_col(c)
                    if 'msisdn' in cn and 'msisdn' not in col_map:
                        col_map['msisdn'] = c
                    elif 'nombre' in cn and 'nombre' not in col_map:
                        col_map['nombre'] = c
                    elif 'paquete' in cn and 'paquete' not in col_map:
                        col_map['paquete'] = c
                    elif 'marca' in cn and 'marca' not in col_map:
                        col_map['marca'] = c
                    elif ('activacion' in cn or 'fecha' in cn) and 'fecha' not in col_map:
                        col_map['fecha'] = c
                    elif 'vendedor' in cn and 'vendedor' not in col_map:
                        col_map['vendedor'] = c

                def g(row, k):
                    c = col_map.get(k)
                    return clean_str(row.get(c, '')) if c else ''

                for idx, (_, r) in enumerate(df.iterrows()):
                    msisdn = g(r, 'msisdn')
                    if not msisdn:
                        continue
                    activaciones_list.append({
                        '_orden': idx,  # posición en la hoja: desempata activaciones del mismo día
                        'msisdn': msisdn,
                        'nombre': fix_mojibake(g(r, 'nombre')) or 'Cliente Bitelit',
                        'paquete': g(r, 'paquete') or 'Sin paquete',
                        'marca': g(r, 'marca') or 'Sin marca',
                        'fecha': self._pd('activaciones', g(r, 'fecha')),
                        'vendedor': g(r, 'vendedor') or 'Directo'
                    })
            except Exception as e:
                self._register_error('activaciones', e)

        pkg_counts, brand_counts, date_counts = {}, {}, {}
        for a in activaciones_list:
            pkg_counts[a['paquete']] = pkg_counts.get(a['paquete'], 0) + 1
            brand_counts[a['marca']] = brand_counts.get(a['marca'], 0) + 1
            if a['fecha']:
                date_counts[a['fecha']] = date_counts.get(a['fecha'], 0) + 1

        return {
            'total': len(activaciones_list),
            'top_packages': sorted([{'name': k, 'count': v} for k, v in pkg_counts.items()], key=lambda x: x['count'], reverse=True)[:6],
            'top_brands': sorted([{'name': k, 'count': v} for k, v in brand_counts.items()], key=lambda x: x['count'], reverse=True)[:6],
            'timeline': sorted([{'date': k, 'count': v} for k, v in date_counts.items()], key=lambda x: x['date'])[-30:],
            # Todas las activaciones (la página filtra por fechas), de la más reciente a la más antigua
            'records': [{k: v for k, v in a.items() if k != '_orden'}
                        for a in sorted(activaciones_list, key=lambda x: (str(x.get('fecha', '')), x['_orden']), reverse=True)]
        }

    # ------------------------------------------------------------------ FINANZAS (Google Sheet de cierre bancario)
    @staticmethod
    def _money(v):
        """Convierte '$347.284,12', '$ 28,000.00', '3231.87' o un número a float."""
        if v is None or v == '':
            return None
        if isinstance(v, (int, float)):
            return float(v)
        s = re.sub(r'[^\d,.\-]', '', str(v))
        if not s or s in ('-', '.', ','):
            return None
        if ',' in s and '.' in s:
            if s.rfind(',') > s.rfind('.'):      # 347.284,12
                s = s.replace('.', '').replace(',', '.')
            else:                                 # 28,000.00
                s = s.replace(',', '')
        elif ',' in s:
            ent, _, dec = s.rpartition(',')
            s = (ent.replace(',', '') + '.' + dec) if len(dec) in (1, 2) else s.replace(',', '')
        elif s.count('.') > 1 or re.match(r'^-?\d{1,3}\.\d{3}$', s):   # 1.160 = mil ciento sesenta
            s = s.replace('.', '')
        try:
            return float(s)
        except ValueError:
            return None

    def _fecha_mov(self, v, anio_ref):
        """'07-10' (día-mes, sin año) -> 'YYYY-MM-DD'. También acepta fechas completas."""
        s = clean_str(v)
        if not s:
            return ''
        m = re.match(r'^(\d{1,2})[-/](\d{1,2})$', s)
        if m:
            d, mes = int(m.group(1)), int(m.group(2))
            anio = anio_ref
            try:
                f = datetime(anio, mes, d)
            except ValueError:
                return ''
            if f > datetime.now() + timedelta(days=1):   # p. ej. '28-12' leído en enero
                f = datetime(anio - 1, mes, d)
            return f.strftime('%Y-%m-%d')
        return parse_date(s, self._dayfirst.get('finanzas', True))

    def _load_finanzas(self):
        cfg = self.config.get('finanzas', {})
        vacio = {'fuente': 'cierre_bancario', 'configurado': bool(cfg.get('sheet_id')), 'movimientos': [],
                 'conceptos': [], 'totales': [], 'capturas': [], 'pasarelas': {}, 'corte': None, 'titulo': ''}
        if not cfg.get('sheet_id'):
            return vacio
        client = self._get_gs_client()
        if not client:
            self._register_error('finanzas', RuntimeError('Sin credenciales de Google Sheets'))
            return vacio
        try:
            doc = client.open_by_key(cfg['sheet_id'])
            self._dayfirst['finanzas'] = not str(getattr(doc, 'locale', '') or '').startswith('en')
            nombres = cfg.get('pestanas') or []
            hojas = [doc.worksheet(n) for n in nombres] if nombres else doc.worksheets()
        except Exception as e:
            self._register_error('finanzas', e)
            return vacio

        out = dict(vacio)
        cols_vistas = {}
        for ws in hojas:
            try:
                vals = ws.get_all_values(value_render_option='UNFORMATTED_VALUE',
                                         date_time_render_option='FORMATTED_STRING')
            except Exception as e:
                self._register_error('finanzas', e)
                continue
            vals = [[clean_str(c) if not isinstance(c, (int, float)) else c for c in r] for r in vals]
            texto = ' '.join(str(c).upper() for r in vals[:40] for c in r if isinstance(c, str))
            cols_vistas[ws.title] = [str(c) for c in (vals[0] if vals else [])]

            if 'HISTORIAL' in ws.title.upper() and 'PASARELA' in ws.title.upper():
                out['pasarelas'] = self._parse_historial(vals)
            elif 'TOTAL INGRESOS' in texto:
                self._parse_resumen(vals, out)
            elif any(self._es_header_mov(r) for r in vals[:5]):
                self._parse_movimientos(vals, out)
        self.headers_found['finanzas'] = cols_vistas
        # más reciente primero; dentro del mismo día se respeta el orden de la hoja (arriba = más nuevo)
        out['movimientos'].sort(key=lambda m: (m['fecha'], -m['_orden']), reverse=True)
        for m in out['movimientos']:
            m.pop('_orden', None)
        if not out['conceptos'] and not out['movimientos']:
            self._register_error('finanzas', RuntimeError(
                'No se encontró la pestaña de resumen (con "TOTAL INGRESOS") ni la de movimientos (con "Fecha" y "Abono").'))
        return out

    @staticmethod
    def _es_header_mov(row):
        n = [normalize_col(c) for c in row if isinstance(c, str)]
        return 'fecha' in n and any(x.startswith('abono') for x in n)

    def _parse_resumen(self, vals, out):
        out['titulo'] = next((c for r in vals[:2] for c in r if isinstance(c, str) and c), '')
        corte = next((c for r in vals[:3] for c in r if isinstance(c, str) and re.search(r'\d{1,2}/\d{1,2}/\d{4}', c)), '')
        m = re.search(r'(\d{1,2})/(\d{1,2})/(\d{4})', corte)
        if m:
            out['corte'] = f'{m.group(3)}-{int(m.group(2)):02d}-{int(m.group(1)):02d}'
        en_conceptos = False
        for r in vals:
            textos = [c for c in r if isinstance(c, str) and c]
            if not textos:
                continue
            etiqueta = textos[0]
            et_up = etiqueta.upper()
            monto = next((self._money(c) for c in r[1:] if self._money(c) is not None), None)
            if et_up == 'CONCEPTO':
                en_conceptos = True
                continue
            if et_up.startswith('TOTAL') or et_up.startswith('SALDO'):
                en_conceptos = False
                if monto is not None:
                    out['totales'].append({'name': etiqueta, 'monto': round(monto, 2)})
                continue
            if en_conceptos and monto is not None:
                out['conceptos'].append({'name': etiqueta, 'monto': round(monto, 2)})

    def _parse_movimientos(self, vals, out):
        h_idx = next(i for i, r in enumerate(vals[:5]) if self._es_header_mov(r))
        header = [normalize_col(c) if isinstance(c, str) else '' for c in vals[h_idx]]
        def col(*nombres):
            for n in nombres:
                for i, h in enumerate(header):
                    if h.startswith(n):
                        return i
            return None
        c_f, c_c = col('fecha'), col('concepto', 'referencia', 'descripcion')
        c_cargo, c_abono, c_saldo = col('cargo'), col('abono'), col('saldo')
        c_cat = col('categoria', 'tipo', 'clasificacion')
        if c_cat is None:  # primera columna sin encabezado después de Saldo
            ult = max(i for i in (c_f, c_c, c_cargo, c_abono, c_saldo) if i is not None)
            c_cat = next((i for i in range(ult + 1, len(header)) if not header[i]), None)
        anio = int(out['corte'][:4]) if out.get('corte') else datetime.now().year
        g = lambda r, i: (r[i] if i is not None and i < len(r) else '')
        for n, r in enumerate(vals[h_idx + 1:]):
            fecha = self._fecha_mov(g(r, c_f), anio)
            cargo, abono = self._money(g(r, c_cargo)), self._money(g(r, c_abono))
            cat = clean_str(g(r, c_cat)).upper()
            if not fecha:
                if cat and abono:   # captura del día de pasarelas (renglones sin fecha arriba de la tabla)
                    out['capturas'].append({'name': cat, 'monto': round(abono, 2)})
                continue
            if cargo is None and abono is None:
                continue
            out['movimientos'].append({
                '_orden': n, 'fecha': fecha,
                'concepto': re.sub(r'\s+', ' ', clean_str(g(r, c_c))),
                'categoria': cat or 'SIN CATEGORÍA',
                'cargo': round(cargo or 0, 2), 'abono': round(abono or 0, 2),
                'saldo': round(self._money(g(r, c_saldo)), 2) if self._money(g(r, c_saldo)) is not None else None,
            })

    def _parse_historial(self, vals):
        if len(vals) < 2:
            return {}
        h = [normalize_col(c) if isinstance(c, str) else '' for c in vals[0]]
        ix = lambda name: h.index(name) if name in h else None
        res = {}
        for r in vals[1:]:
            nombre = clean_str(r[ix('pasarela')]) if ix('pasarela') is not None else ''
            if not nombre:
                continue
            g = lambda k: r[ix(k)] if ix(k) is not None and ix(k) < len(r) else ''
            res[nombre] = {   # el último renglón de cada pasarela gana
                'fecha': clean_str(g('fecha')), 'hora': clean_str(g('hora')),
                'saldo_a_favor': self._money(g('saldo a favor')), 'disponible': self._money(g('disponible')),
                'proxima_dispersion': self._money(g('proxima dispersion')),
                'fecha_dispersion': clean_str(g('fecha dispersion')), 'saldo_total': self._money(g('saldo total')),
                'ultima_dispersion': self._money(g('ultima dispersion')),
                'fecha_ultima_dispersion': clean_str(g('fecha ultima dispersion')),
                'neto_ayer': self._money(g('ventas ayer neto')), 'bruto_ayer': self._money(g('ventas ayer bruto')),
                'transacciones_ayer': self._money(g('transacciones ayer')),
                'error': clean_str(g('error')),
            }
        return res

    def _build_overview(self, tickets, procesos, activaciones, finanzas):
        return {
            'kpis': {
                'total_tickets': tickets.get('total', 0),
                'tickets_cerrados_pct': tickets.get('tasa_resolucion', 0),
                'total_procesos': procesos.get('total', 0),
                'sim_swaps': procesos.get('sim_swap_count', 0),
                'cambios_dn': procesos.get('cambio_dn_count', 0),
                'total_activaciones': activaciones.get('total', 0),
                'ingresos_totales': sum(m['abono'] for m in finanzas.get('movimientos', [])),
            },
            'top_motivos': tickets.get('motivos', [])[:5],
            'top_tramites': procesos.get('by_type', [])[:5],
            'top_bancos': finanzas.get('conceptos', [])
        }

if __name__ == '__main__':
    # Prueba rápida desde la raíz del proyecto:  python -m bitelit.data_loader
    loader = BitelitDataLoader()
    data = loader.load_all_data(force_reload=True)
    print("Overview KPIs:", data['overview']['kpis'])
    print("Tickets count:", data['tickets']['total'])
    print("Procesos count:", data['procesos']['total'])
    print("Activaciones count:", data['activaciones']['total'])
    print("Finanzas movimientos:", len(data['finanzas']['movimientos']), "totales:", data['finanzas']['totales'])
